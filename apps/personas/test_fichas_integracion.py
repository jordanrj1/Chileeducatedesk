"""F24 T11: la fusión, la exportación de datos personales, la importación y el rol en la actividad conocen las fichas por tipo,
y los datos del hogar no salen en el Excel territorial ni en el buscador de personas. Todos los datos son ficticios."""
import importlib
import io

import pytest
from django.apps import apps as registro
from openpyxl import load_workbook

from apps.actividades.models import Participacion
from apps.importaciones.models import Lote
from apps.lugares.models import Lugar
from apps.nucleo.fusion import fusionar
from apps.nucleo.models import EstadoVoluntario, Parentesco, PerfilPersona, RolParticipacion, TipoInvitado, TipoLugar, TramoEtario
from apps.personas.models import (
    Contacto, FamiliarVinculado, FichaInvitado, FichaVecino, FichaVoluntario, IntegranteHogar, Persona, PersonaDireccion, PersonaFormacion,
)
from apps.territorio.models import Direccion

pytestmark = pytest.mark.django_db
rol_perfil = importlib.import_module("apps.nucleo.migrations.0015_rol_perfil")


def _par(codigo):
    return Parentesco.objects.get(codigo=codigo)


def _perfiles(p):
    return set(p.perfiles.values_list("codigo", flat=True))


def _textos(libro):
    return " | ".join(str(c) for h in libro for fila in h.iter_rows(values_only=True) for c in fila if c is not None)


# --- 1. Fusión ---------------------------------------------------------------------------------
def test_fusion_completa_las_fichas_sin_pisar_y_no_suma_el_mismo_hogar(crear_persona):
    origen, destino = crear_persona(), crear_persona()
    fo = FichaVecino.objects.create(persona=origen, jefatura_hogar=True, integrantes_hogar=5)
    fd = FichaVecino.objects.create(persona=destino, integrantes_hogar=3)
    IntegranteHogar.objects.create(ficha=fo, parentesco=_par("hijo_o_hija"), cantidad=2)  # repite la fila del destino
    IntegranteHogar.objects.create(ficha=fo, parentesco=_par("nieto_o_nieta"), vive_en_hogar=False)
    IntegranteHogar.objects.create(ficha=fd, parentesco=_par("hijo_o_hija"), cantidad=1)
    FichaVoluntario.objects.create(persona=origen, horas_semanales=4, estado=EstadoVoluntario.objects.get(codigo="activo"))
    FichaVoluntario.objects.create(persona=destino, motivacion="Quiere apoyar talleres.")
    FichaInvitado.objects.create(persona=origen, cargo="Directora ficticia")  # solo el origen: la mueve el bucle genérico

    _, conservados = fusionar(origen, destino, None, "misma persona")

    fd.refresh_from_db()
    assert fd.jefatura_hogar is True and fd.integrantes_hogar == 3  # completa lo vacío, no pisa
    assert sorted((str(i.parentesco), i.cantidad) for i in fd.integrantes.all()) == [("Hijo o hija", 1), ("Nieto o nieta", 1)]
    assert list(fo.integrantes.values_list("cantidad", flat=True)) == [2]  # la fila repetida se queda en el origen
    vol = FichaVoluntario.objects.get(persona=destino)
    assert vol.horas_semanales == 4 and vol.estado.codigo == "activo" and vol.motivacion == "Quiere apoyar talleres."
    assert FichaInvitado.objects.get().persona == destino
    assert FichaVecino.objects.filter(persona=origen).exists() and any("ficha de vecino" in c for c in conservados)


def test_fusion_no_deja_a_la_persona_como_su_propio_familiar_ni_repetida(crear_persona):
    origen, destino, otro, tercero = crear_persona(), crear_persona(), crear_persona(), crear_persona()
    FamiliarVinculado.objects.create(persona=origen, familiar=destino, parentesco=_par("hermano_o_hermana"))
    FamiliarVinculado.objects.create(persona=origen, familiar=otro, parentesco=_par("hijo_o_hija"))
    FamiliarVinculado.objects.create(persona=destino, familiar=otro, parentesco=_par("hijo_o_hija"))
    FamiliarVinculado.objects.create(persona=origen, familiar=tercero, parentesco=_par("nieto_o_nieta"))

    fusionar(origen, destino, None, "misma persona")

    assert set(destino.familiares.values_list("familiar_id", flat=True)) == {otro.pk, tercero.pk}
    assert set(origen.familiares.values_list("familiar_id", flat=True)) == {destino.pk, otro.pk}  # se conservan en el origen
    assert not FamiliarVinculado.objects.filter(persona=destino, familiar=destino).exists()


# --- 2. Exportación de datos personales ---------------------------------------------------------
def test_exportacion_de_datos_personales_lleva_las_fichas(cg, crear_persona, universidad, comuna):
    p = crear_persona(sexo="mujer")
    p.fijar_edad_aproximada(70)
    p.save()
    nieto = crear_persona(nombres="Tomás")
    fv = FichaVecino.objects.create(persona=p, jefatura_hogar=True, integrantes_hogar=3)
    IntegranteHogar.objects.create(ficha=fv, parentesco=_par("nieto_o_nieta"), tramo=TramoEtario.objects.order_by("orden").first(), cantidad=2, vive_en_hogar=False)
    FamiliarVinculado.objects.create(persona=p, familiar=nieto, parentesco=_par("nieto_o_nieta"))
    FichaInvitado.objects.create(persona=p, tipo=TipoInvitado.objects.get(codigo="autoridad"), cargo="Presidenta de junta ficticia")
    FichaVoluntario.objects.create(persona=p, horas_semanales=6)
    sede = Lugar.objects.create(nombre="Sede Centro Ficticia", tipo=TipoLugar.objects.get(codigo="sala"),
                                direccion=Direccion.objects.create(calle="Calle Inventada", numero="1", comuna=comuna), organizacion_propietaria=universidad)
    PersonaFormacion.objects.create(persona=p, organizacion=universidad, sede_institucion=sede)

    r = cg.get(f"/personas/{p.pk}/datos-personales")
    libro = load_workbook(io.BytesIO(r.content))
    assert {"Ficha de vecino", "Ficha de invitado", "Ficha de voluntario"} <= set(libro.sheetnames)
    datos = dict(libro["Datos"].iter_rows(min_row=2, values_only=True))
    assert datos["Sexo"] == "Mujer" and datos["Fecha de nacimiento aproximada"] == "Sí"
    vecino = _textos([libro["Ficha de vecino"]])
    assert "Jefatura de hogar | Sí" in vecino and "Nieto o nieta" in vecino and " · no" in vecino and "Tomás" in vecino
    assert "Presidenta de junta ficticia" in _textos([libro["Ficha de invitado"]]) and "6" in _textos([libro["Ficha de voluntario"]])
    assert "Sede Centro Ficticia" in _textos([libro["Formación"]])
    # sin fichas no hay hojas vacías
    otra = crear_persona()
    libro = load_workbook(io.BytesIO(cg.get(f"/personas/{otra.pk}/datos-personales").content))
    assert not {"Ficha de vecino", "Ficha de invitado", "Ficha de voluntario"} & set(libro.sheetnames)


# --- 3. Importación ----------------------------------------------------------------------------
def test_importacion_con_columna_tipo_agrega_el_perfil_y_avisa_si_no_existe(cg, crear_persona):
    existente = crear_persona(nombres="Rosa", apellido_paterno="Fuentes")
    existente.perfiles.add(PerfilPersona.objects.get(codigo="estudiante"))
    Contacto.objects.create(persona=existente, tipo_id=registro.get_model("nucleo", "TipoContacto").objects.get(codigo="correo").pk, valor="rosa@ejemplo.invalid")
    texto = ("Nombre\tCorreo\tTipo\n"
             "Elena Soto\telena@ejemplo.invalid\tVOLUNTARIO o voluntária\n"
             "Rosa Fuentes\trosa@ejemplo.invalid\tvecino_o_vecina_de_la_comunidad\n"
             "Pedro Rojas\tpedro@ejemplo.invalid\tAstronauta\n")
    cg.post("/importaciones/nueva", {"tipo": "personas", "texto": texto})
    lote = Lote.objects.get()
    assert [f["estado"] for f in lote.filas] == ["nuevo", "existente", "nuevo"]  # un perfil desconocido no es error
    assert any("Astronauta" in a for a in lote.filas[2]["avisos"]) and not lote.filas[0]["avisos"]
    cg.post(f"/importaciones/{lote.pk}/revisar", {"accion": "confirmar"})
    assert _perfiles(Persona.objects.get(nombres="Elena")) == {"voluntario_o_voluntaria"}
    assert _perfiles(existente) == {"estudiante", "vecino_o_vecina_de_la_comunidad"}  # se suma, no quita
    assert _perfiles(Persona.objects.get(nombres="Pedro")) == set()
    cg.post(f"/importaciones/{lote.pk}/revertir")
    assert _perfiles(existente) == {"estudiante", "vecino_o_vecina_de_la_comunidad"}


# --- 4. Rol → perfil ---------------------------------------------------------------------------
def test_perfil_inicial_de_los_roles_y_su_edicion_en_catalogos(cg, rol):
    perfiles = dict(RolParticipacion.objects.values_list("codigo", "perfil__codigo"))
    assert perfiles["practicante"] == "estudiante" and perfiles["voluntario"] == "voluntario_o_voluntaria"
    assert perfiles["tutor"] == "docente_o_academico" and perfiles["contacto_institucional"] == "funcionario_institucional"
    assert perfiles["coordinador"] is None and perfiles["participante"] is None
    # repetir la carga no pisa lo editado
    RolParticipacion.objects.filter(codigo="voluntario").update(perfil=PerfilPersona.objects.get(codigo="dirigente_social"))
    rol_perfil.asignar(registro, None)
    assert rol("voluntario").perfil.codigo == "dirigente_social"
    assert 'name="perfil"' in cg.get(f"/administracion/catalogos/rol-participacion/{rol('tallerista').pk}/editar").content.decode()


def test_sumar_a_la_actividad_agrega_el_perfil_del_rol(cg, actividad, rol, estado_part, crear_persona):
    cg.post(f"/actividades/{actividad.pk}/participantes/rapido", {"nombres": "Lucía", "apellido_paterno": "Vera", "rol": rol("voluntario").pk, "confirmar_duplicado": "1"})
    assert _perfiles(Persona.objects.get(nombres="Lucía")) == {"voluntario_o_voluntaria"}
    p = crear_persona()
    p.perfiles.add(PerfilPersona.objects.get(codigo="estudiante"))
    r = cg.post(f"/actividades/{actividad.pk}/participantes/nuevo", {"persona": p.pk, "rol": rol("tallerista").pk, "estado": estado_part("confirmado").pk})
    assert r.status_code == 302 and Participacion.objects.filter(persona=p).exists()
    assert _perfiles(p) == {"estudiante", "tallerista_o_profesional"}


# --- 5. Privacidad -----------------------------------------------------------------------------
def test_los_datos_del_hogar_no_salen_en_el_excel_territorial_ni_en_el_buscador(cg, gestor, comuna, crear_persona):
    from apps.territorio.exportar import libro_territorial

    vecina = crear_persona(nombres="Vecina", apellido_paterno="Ficticia")
    PersonaDireccion.objects.create(persona=vecina, direccion=Direccion.objects.create(calle="Pasaje Inventado", numero="7", comuna=comuna), tipo="residencia")
    fv = FichaVecino.objects.create(persona=vecina, jefatura_hogar=True, integrantes_hogar=4)
    IntegranteHogar.objects.create(ficha=fv, parentesco=_par("nieto_o_nieta"), cantidad=3)
    FamiliarVinculado.objects.create(persona=vecina, familiar=crear_persona(), parentesco=_par("abuelo_o_abuela"))

    excel = _textos(libro_territorial(gestor, comuna=comuna))
    assert "Vecina Ficticia" in excel  # la persona sí sale, como residente
    opciones = cg.get("/personas/opciones?q=Vecina").content.decode()
    assert "Vecina" in opciones
    for texto in (excel, opciones):
        for dato in ("Nieto o nieta", "Abuelo o abuela", "Jefatura", "jefe o jefa", "Integrantes del hogar", "integrantes"):
            assert dato not in texto
