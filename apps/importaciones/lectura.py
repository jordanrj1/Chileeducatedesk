"""Lectura de planillas (CSV, TXT con tabulaciones, XLSX) y texto pegado. Las celdas se leen como texto."""
import csv
import io
import re
import unicodedata

from openpyxl import load_workbook

MAX_FILAS = 5000

SINONIMOS = {
    "nombre_completo": {"nombre completo", "nombre y apellido", "nombres y apellidos", "estudiante", "participante", "alumno", "alumna", "persona"},
    "nombres": {"nombres", "nombre", "primer nombre"},
    "apellido_paterno": {"apellido paterno", "paterno", "primer apellido", "ap paterno", "apellido 1"},
    "apellido_materno": {"apellido materno", "materno", "segundo apellido", "ap materno", "apellido 2"},
    "apellidos": {"apellidos", "apellido"},
    "rut": {"rut", "run", "rut/run", "identificador", "cedula", "cédula"},
    "correo": {"correo", "correo electronico", "email", "e-mail", "mail", "correo institucional", "correo personal"},
    "telefono": {"telefono", "teléfono", "fono", "celular", "movil", "móvil", "telefono movil", "whatsapp", "contacto"},
    "rol": {"rol", "cargo", "tipo de participante"},
    "carrera": {"carrera", "programa", "especialidad"},
    "institucion": {"institucion", "institución", "universidad", "casa de estudios", "establecimiento"},
    "profesion": {"profesion", "profesión", "oficio", "ocupacion", "ocupación"},
    "observaciones": {"observaciones", "observacion", "observación", "comentarios", "notas"},
    # organizaciones / lugares
    "tipo": {"tipo", "tipo de organizacion", "tipo de organización", "tipo de lugar", "categoria", "categoría"},
    "comuna": {"comuna"},
    "calle": {"calle", "direccion", "dirección", "domicilio"},
    "numero": {"numero", "número", "n°", "nro", "n"},
    "capacidad": {"capacidad", "aforo", "cupos"},
    "condiciones": {"condiciones", "equipamiento"},
    "sitio_web": {"sitio web", "web", "pagina web", "página web"},
}
INVERSO = {}
for canon, nombres in SINONIMOS.items():
    for n in nombres:
        INVERSO.setdefault(n, canon)


def _sin_tildes(t):
    return "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")


def canon_cabecera(h):
    clave = re.sub(r"\s+", " ", _sin_tildes(str(h or "")).strip().lower().replace("_", " "))
    return INVERSO.get(clave) or INVERSO.get(str(h or "").strip().lower())


def _texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _decodificar(datos: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return datos.decode(enc)
        except UnicodeDecodeError:
            continue
    return datos.decode("utf-8", errors="replace")


def _filas_de_texto(texto: str):
    muestra = texto[:4096]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
        if muestra.count("\t") > muestra.count(","):
            dialecto = csv.excel_tab
        elif muestra.count(";") > muestra.count(","):
            dialecto.delimiter = ";"
    return list(csv.reader(io.StringIO(texto), dialecto))


def _filas_xlsx(datos: bytes):
    libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    hoja = libro.active
    return [[_texto(c) for c in fila] for fila in hoja.iter_rows(values_only=True)]


class ErrorLectura(Exception):
    pass


COLUMNAS_POR_TIPO = {
    "personas": [("nombre_completo", "Nombre completo"), ("nombres", "Nombres"), ("apellido_paterno", "Apellido paterno"), ("apellido_materno", "Apellido materno"),
                 ("apellidos", "Apellidos (ambos)"), ("rut", "RUT"), ("correo", "Correo"), ("telefono", "Teléfono"), ("institucion", "Institución"),
                 ("carrera", "Carrera"), ("profesion", "Profesión"), ("observaciones", "Observaciones")],
    "organizaciones": [("nombres", "Nombre"), ("tipo", "Tipo"), ("rut", "RUT"), ("correo", "Correo"), ("telefono", "Teléfono"), ("comuna", "Comuna"), ("calle", "Dirección"), ("sitio_web", "Sitio web")],
    "lugares": [("nombres", "Nombre"), ("tipo", "Tipo"), ("calle", "Calle"), ("numero", "Número"), ("comuna", "Comuna"), ("capacidad", "Capacidad"), ("condiciones", "Condiciones")],
}
COLUMNAS_POR_TIPO["listado_actividad"] = COLUMNAS_POR_TIPO["personas"]


def _es_fila_de_total(fila):
    """IG-E3: «Total», «Subtotal», «TOTAL ALUMNOS»… al comienzo de la fila."""
    primera = next((_texto(c) for c in fila if _texto(c)), "")
    return _sin_tildes(primera).lower().startswith(("total", "subtotal"))


def leer_tabla(contenido: bytes | None = None, nombre: str = "", texto: str = ""):
    """(encabezados, filas, filas_de_total_ignoradas). Las celdas siempre como texto."""
    try:
        if contenido is not None:
            filas = _filas_xlsx(contenido) if nombre.lower().endswith(".xlsx") else _filas_de_texto(_decodificar(contenido))
        else:
            filas = _filas_de_texto(texto)
    except Exception as e:  # archivo corrupto o formato inesperado
        raise ErrorLectura(f"No se pudo leer el archivo: {e}")
    filas = [[_texto(c) for c in f] for f in filas if any(_texto(c) for c in f)]
    if len(filas) < 2:
        raise ErrorLectura("El archivo no tiene datos: se espera una fila de encabezados y al menos una fila de datos.")
    encabezados, datos = filas[0], filas[1:]
    totales = [f for f in datos if _es_fila_de_total(f)]
    datos = [f for f in datos if not _es_fila_de_total(f)]
    if len(datos) > MAX_FILAS:
        raise ErrorLectura(f"Demasiadas filas ({len(datos)}). El máximo por carga es {MAX_FILAS}.")
    if not datos:
        raise ErrorLectura("El archivo no tiene filas de datos (solo encabezados o totales).")
    return encabezados, datos, len(totales)


def mapeo_sugerido(encabezados):
    return [canon_cabecera(h) or "" for h in encabezados]


def aplicar_mapeo(filas, mapeo):
    resultado = []
    for fila in filas:
        d = {}
        for c, v in zip(mapeo, fila):
            if c and v and not d.get(c):
                d[c] = v
        resultado.append(d)
    return resultado


def leer(contenido: bytes | None = None, nombre: str = "", texto: str = ""):
    """Devuelve (filas, columnas_ignoradas). Cada fila es un dict con claves canónicas."""
    encabezados, datos, _ = leer_tabla(contenido, nombre, texto)
    mapeo = mapeo_sugerido(encabezados)
    if not any(mapeo):
        raise ErrorLectura("No se reconocieron los encabezados. Usa columnas como: nombre, apellidos, rut, correo, teléfono.")
    ignoradas = [str(h) for h, c in zip(encabezados, mapeo) if not c and str(h).strip()]
    return aplicar_mapeo(datos, mapeo), ignoradas
