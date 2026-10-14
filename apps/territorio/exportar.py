"""Excel del resumen territorial (F23): lo que hay y lo que se ha hecho en una comuna o en un sector, para la Dirección y para un fondo.

Hojas, en este orden: Resumen, Sectores, Comunidades, Organizaciones, Personas residentes, Actividades, Alcance, más el aviso de
confidencialidad. Cada hoja aparece solo si quien exporta puede ver ese dato en la ficha (los mismos permisos de las pestañas).

Datos personales (Ley 21.719, RNF-02):
- «Personas residentes» lleva nombre, perfiles y comunidad. RUT, correo y teléfono salen solo con `personas.exportar_persona`:
  sin ese permiso las columnas no existen (no van vacías).
- Los menores de edad residentes (`Persona.es_menor`) no se listan por nombre: una fila final dice cuántos son (V#2). Quien no tiene
  fecha de nacimiento se lista, porque su edad se desconoce.
- Pertenecer a una comunidad indígena revela pueblo originario (dato sensible): esas pertenencias no aparecen en la celda «Comunidad»
  de la persona; una fila final dice cuántas se omitieron. La hoja «Comunidades» puede listar la comunidad y su tipo (no lleva nombres de personas).
- Nunca salen datos sensibles (perfil sensible, apoderados, contactos de emergencia, verificaciones, género, salud…) ni el desglose por género del alcance.
- Todo texto pasa por `_seguro` (sin inyección de fórmulas). Los números y las fechas se escriben como tales para que se puedan sumar y ordenar.

Número fijo de consultas, sin importar cuántos registros haya.
"""
from datetime import date
from decimal import Decimal

from django.db.models import Count, Prefetch, Q
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.actividades.alcance import alcance
from apps.actividades.models import AgendaActividad
from apps.comunidades.models import Comunidad, ComunidadMiembro
from apps.nucleo.exportacion import _seguro, agregar_aviso
from apps.personas.models import Contacto

from . import consultas
from .resumen import personas_residentes, resumen_territorio

SECTOR_CON_PADRES = "sector__sector_padre__sector_padre"  # el nombre de un sector incluye el de su padre y su abuelo
CONTACTOS = ("correo", "telefono_movil", "telefono_fijo")  # de más a menos preferido cuando ninguno está marcado como principal
TELEFONOS = ("telefono_movil", "telefono_fijo")
FORMATO_FECHA = "DD-MM-YYYY"
ANCHO_MAXIMO = 60

# Hoja -> permiso que hace falta para verla en pantalla; sin permiso la hoja no se exporta
PERMISO_HOJA = {
    "Comunidades": "comunidades.view_comunidad",
    "Organizaciones": "organizaciones.view_organizacion",
    "Personas residentes": "personas.view_persona",
    "Actividades": "actividades.view_actividad",
    "Alcance": "actividades.view_participacion",
}
PERMISO_RUT_Y_CONTACTOS = "personas.exportar_persona"
DETALLE_RUT_Y_CONTACTOS = "con RUT y contactos"
NOTA_MENORES = "Menores de edad residentes"
NOTA_INDIGENAS = "Pertenencias a comunidades indígenas omitidas"
TIPO_INDIGENA = "indigena"  # Comunidad.TIPOS


def _celda(valor):
    """Texto neutralizado con `_seguro`; las cifras y las fechas se dejan como tales (no pueden ser fórmulas)."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, (int, float, Decimal, date)):
        return valor
    return _seguro(valor)


def _agregar_hoja(libro, titulo, encabezados, filas):
    hoja = libro.create_sheet(titulo)
    hoja.append([_seguro(e) for e in encabezados])
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        hoja.append([_celda(v) for v in fila])
        for celda in hoja[hoja.max_row]:
            if isinstance(celda.value, date):
                celda.number_format = FORMATO_FECHA
    return hoja


def _ajustar_anchos(hoja):
    """Columnas legibles sin que quien abre el archivo tenga que ensancharlas a mano."""
    for columna in hoja.columns:
        hoja.column_dimensions[columna[0].column_letter].width = min(max(len(str(c.value or "")) for c in columna) + 2, ANCHO_MAXIMO)


def _prefetch_contactos():
    """Contactos vigentes de los tipos que se exportan, en `contactos_exportables` (sin consultas por fila)."""
    vigentes = Contacto.objects.filter(vigente=True, tipo__codigo__in=CONTACTOS).select_related("tipo")
    return Prefetch("contactos", queryset=vigentes, to_attr="contactos_exportables")


def _contacto(sujeto, codigos):
    """El contacto marcado como principal entre los `codigos`; si ninguno lo está, el de tipo más preferido."""
    candidatos = [c for c in sujeto.contactos_exportables if c.tipo.codigo in codigos]
    if not candidatos:
        return ""
    return min(candidatos, key=lambda c: (not c.principal, codigos.index(c.tipo.codigo), c.pk)).valor


def _sector(direccion):
    return str(direccion.sector) if direccion and direccion.sector_id else ""


def _hoja_resumen(libro, r, titulo, con_comuna, total_alcanzado):
    filas = [("Territorio", titulo)]
    if con_comuna:
        filas.append(("Comuna", con_comuna))
    filas += [
        ("Comunidades activas", r["comunidades"]),
        ("Organizaciones activas", r["organizaciones"]),
        ("Personas residentes", r["personas"]),
        ("Lugares activos", r["lugares"]),
        ("Actividades de los últimos 12 meses", r["actividades_12m"]),
        ("Actividades en total", r["actividades_total"]),
    ]
    if total_alcanzado is not None:
        filas.append(("Total alcanzado", total_alcanzado))
    filas += [(f"Organizaciones activas · {tipo.nombre}", total) for tipo, total, _ in r["organizaciones_por_tipo"]]
    _agregar_hoja(libro, "Resumen", ["Indicador", "Valor"], filas)


def _hoja_sectores(libro, r):
    filas = [
        (
            str(s["sector"]) if s["sector"] else "Sin sector",
            s["sector"].get_tipo_display() if s["sector"] else "",
            s["comunidades"], s["organizaciones"], s["lugares"], s["personas"], s["actividades"],
        )
        for s in r["sectores"]
    ]
    _agregar_hoja(libro, "Sectores", ["Sector", "Tipo", "Comunidades", "Organizaciones", "Lugares", "Personas residentes", "Actividades"], filas)


def _hoja_comunidades(libro, comuna, sectores):
    filtro = {"sector_id__in": sectores} if sectores else {"comuna": comuna}
    comunidades = (
        Comunidad.objects.filter(activo=True, **filtro)
        .select_related(SECTOR_CON_PADRES, "organizacion")
        .annotate(n_integrantes=Count("miembros", filter=Q(miembros__vigente=True)))
        .order_by("nombre", "pk")
    )
    filas = [
        (c.nombre, c.get_tipo_display(), str(c.sector) if c.sector_id else "", c.organizacion.nombre if c.organizacion_id else "", c.n_integrantes)
        for c in comunidades
    ]
    _agregar_hoja(libro, "Comunidades", ["Comunidad", "Tipo", "Sector", "Organización que la representa", "Integrantes vigentes"], filas)


def _hoja_organizaciones(libro, comuna, sectores):
    """Activas primero y, al final, las desactivadas (columna «Activa»): lo que tiene historia sigue a la vista."""
    organizaciones = (
        consultas.organizaciones_de(comuna, sectores)
        .select_related(f"direccion__{SECTOR_CON_PADRES}")
        .prefetch_related(_prefetch_contactos())
        .order_by("-activo", "tipo__orden", "tipo__nombre", "nombre", "pk")
    )
    filas = [(o.nombre, o.tipo.nombre, _sector(o.direccion), _contacto(o, CONTACTOS), o.activo) for o in organizaciones]
    _agregar_hoja(libro, "Organizaciones", ["Nombre", "Tipo", "Sector", "Correo o teléfono principal", "Activa"], filas)


def _hoja_personas(libro, comuna, sectores, con_rut_y_contactos):
    membresias = ComunidadMiembro.objects.filter(vigente=True, comunidad__activo=True).select_related("comunidad").order_by("comunidad__nombre")
    prefetch = ["perfiles", Prefetch("comunidades", queryset=membresias, to_attr="membresias")]
    if con_rut_y_contactos:
        prefetch.append(_prefetch_contactos())
    encabezados = ["Nombre", "Perfiles", "Comunidad"] + (["RUT", "Correo", "Teléfono"] if con_rut_y_contactos else [])
    filas, menores, indigenas = [], 0, 0
    for p in personas_residentes(comuna, sectores).prefetch_related(*prefetch):
        if p.es_menor:  # menores: nunca por nombre; solo se cuentan (V#2). Sin fecha de nacimiento no se sabe la edad: se listan
            menores += 1
            continue
        comunidades = [m.comunidad.nombre for m in p.membresias if m.comunidad.tipo != TIPO_INDIGENA]
        indigenas += len(p.membresias) - len(comunidades)  # pertenecer a una comunidad indígena revela pueblo originario: dato sensible
        fila = [p.nombre_completo, ", ".join(x.nombre for x in p.perfiles.all()), ", ".join(comunidades)]
        if con_rut_y_contactos:
            fila += [p.rut, _contacto(p, ("correo",)), _contacto(p, TELEFONOS)]
        filas.append(fila)
    if menores:
        filas.append([f"{NOTA_MENORES}: {menores} (no se listan por protección de datos; pendiente V#2)"])
    if indigenas:
        filas.append([f"{NOTA_INDIGENAS}: {indigenas} (dato sensible, Ley 21.719)"])
    _agregar_hoja(libro, "Personas residentes", encabezados, filas)


def _actividades(comuna, sectores):
    """Actividades del territorio con su programación vigente en `agenda_exportar`; la fecha más nueva primero y las sin fecha al final."""
    vigente = AgendaActividad.objects.filter(vigente=True).select_related("lugar")
    lista = list(consultas.actividades_de(comuna, sectores).prefetch_related(Prefetch("agendas", queryset=vigente, to_attr="agenda_exportar")))
    con_fecha = sorted((a for a in lista if a.agenda_exportar), key=lambda a: (a.agenda_exportar[0].fecha, a.pk), reverse=True)
    return con_fecha + [a for a in lista if not a.agenda_exportar]


def _hoja_actividades(libro, actividades):
    filas = []
    for a in actividades:
        agenda = a.agenda_exportar[0] if a.agenda_exportar else None
        filas.append((a.nombre, a.estado.nombre, agenda.fecha if agenda else None, agenda.lugar.nombre if agenda and agenda.lugar_id else ""))
    _agregar_hoja(libro, "Actividades", ["Nombre", "Estado", "Fecha", "Lugar"], filas)


def _hoja_alcance(libro, a):
    """El alcance de las actividades como indicador / valor; instituciones, comunidades y roles como listas aparte. Sin desglose por género."""
    hoja = _agregar_hoja(libro, "Alcance", ["Indicador", "Valor"], [
        ("Actividades", a["actividades"]),
        ("Personas con participación registrada", a["personas"]),
        ("Equipo ejecutor", a["equipo"]),
        ("Asistentes registrados", a["asistentes"]),
        ("Asistieron (personas registradas)", a["asistieron"]),
        ("Asistentes sin registro", a["sin_registro"]),
        ("Total alcanzado", a["total_alcanzado"]),
        ("Horas de participación", a["horas"]),
        ("Horas de trabajo del equipo", a["horas_equipo"]),
        ("Sesiones", a["sesiones"]),
        *[(f"Edad {tramo} (registradas)", n) for tramo, n in a["edades_registradas"]],
        *[(f"Edad {tramo} (sin registro)", n) for tramo, n in a["sin_registro_edades"]],
    ])
    for titulo, columna, pares in (
        ("Instituciones de procedencia del equipo", "Participaciones", a["instituciones"]),
        ("Comunidades de procedencia", "Participaciones", a["comunidades"]),
        ("Participaciones por rol", "Participaciones", a["roles"]),
    ):
        hoja.append([])
        hoja.append([_seguro(titulo), _seguro(columna)])
        for objeto, n in pares or [("Sin datos", None)]:
            hoja.append([_celda(str(objeto)), _celda(n)])


def libro_territorial(usuario, comuna=None, sectores=None, titulo=""):
    """Libro de Excel con el resumen de una comuna (`comuna`) o de un sector y los que contiene (`sectores`: sus ids; `comuna`: la suya).

    Lo que lleva depende de lo que `usuario` puede ver (ver el módulo). Devuelve el `Workbook` listo para guardar.
    """
    titulo = titulo or (str(comuna) if comuna else "")
    visibles = {hoja: usuario.has_perm(permiso) for hoja, permiso in PERMISO_HOJA.items()}
    r = resumen_territorio(comuna=comuna, sectores=sectores)
    actividades = _actividades(comuna, sectores) if visibles["Actividades"] or visibles["Alcance"] else []
    datos_alcance = alcance(actividades) if visibles["Alcance"] else None

    libro = Workbook()
    libro.remove(libro.active)
    _hoja_resumen(libro, r, titulo, comuna.nombre if sectores and comuna else "", datos_alcance["total_alcanzado"] if datos_alcance else None)
    _hoja_sectores(libro, r)
    if visibles["Comunidades"]:
        _hoja_comunidades(libro, comuna, sectores)
    if visibles["Organizaciones"]:
        _hoja_organizaciones(libro, comuna, sectores)
    if visibles["Personas residentes"]:
        _hoja_personas(libro, comuna, sectores, usuario.has_perm(PERMISO_RUT_Y_CONTACTOS))
    if visibles["Actividades"]:
        _hoja_actividades(libro, actividades)
    if datos_alcance:
        _hoja_alcance(libro, datos_alcance)
    for hoja in libro:
        _ajustar_anchos(hoja)
    agregar_aviso(libro, usuario)
    return libro


def detalle_de(usuario, libro):
    """Texto para la auditoría: qué hojas salieron y si llevaron RUT y contactos."""
    hojas = [h for h in libro.sheetnames if h != "Aviso"]
    con = "Personas residentes" in hojas and usuario.has_perm(PERMISO_RUT_Y_CONTACTOS)
    return f"Hojas: {', '.join(hojas)}" + (f"; {DETALLE_RUT_Y_CONTACTOS}" if con else "")
