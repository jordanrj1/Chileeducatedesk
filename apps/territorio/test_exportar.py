"""Exportar el resumen territorial a Excel (F23): hojas, permisos, datos personales y auditoría. Datos ficticios."""
from datetime import timedelta
from io import BytesIO

import pytest
from django.contrib.auth.models import Permission, User
from django.db import connection
from django.test import Client
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.actividades.models import Participacion, SesionActividad
from apps.actividades.servicios import programar
from apps.comunidades.models import Comunidad, ComunidadMiembro
from apps.nucleo.models import Auditoria, EstadoParticipacion, PerfilPersona, RolParticipacion, TipoContacto, TipoOrganizacion
from apps.nucleo.utilidades import calcular_dv, formatear_rut
from apps.organizaciones.models import Organizacion
from apps.personas.models import Contacto, ContactoEmergencia, PerfilSensible, PersonaDireccion
from apps.territorio.models import Direccion, Sector

pytestmark = pytest.mark.django_db

HOY = timezone.localdate()
HOJAS = ["Resumen", "Sectores", "Comunidades", "Organizaciones", "Personas residentes", "Actividades", "Alcance"]
VISTAS = ["view_comuna", "view_sector", "view_comunidad", "view_organizacion", "view_lugar", "view_actividad", "view_participacion", "view_persona"]
RUT = 12345678
CORREO = "rosa.ficticia@ejemplo.invalid"
TELEFONO = "+56911112222"


# ---------- ayudas ----------

def _usuario(nombre, *permisos):
    u = User.objects.create_user(nombre, email=f"{nombre}@ejemplo.invalid", password="clave-ficticia-123")
    for codigo in permisos:
        u.user_permissions.add(Permission.objects.get(codename=codigo))
    return u


def _cliente(usuario):
    c = Client()
    c.force_login(usuario)
    return c


def _dir(comuna, sector=None):
    return Direccion.objects.create(calle="Calle Inventada", numero="1", comuna=comuna, sector=sector)


def _libro(respuesta):
    assert respuesta.status_code == 200, respuesta.status_code
    return load_workbook(BytesIO(respuesta.content))


def _filas(libro, hoja):
    return [list(f) for f in libro[hoja].iter_rows(values_only=True)]


def _celdas(libro):
    return [str(v) for h in libro for f in h.iter_rows(values_only=True) for v in f if v is not None]


def _participa(actividad, persona, rol="participante", asistencia=True):
    return Participacion.objects.create(
        actividad=actividad, persona=persona, rol=RolParticipacion.objects.get(codigo=rol),
        estado=EstadoParticipacion.objects.get(codigo="asistio" if asistencia else "confirmado"), asistencia=asistencia,
    )


@pytest.fixture
def mundo(comuna, lugar, crear_actividad, crear_persona, gestor):
    """Una comuna con sector y etapa, una comunidad, organizaciones, una residente con datos personales y sensibles, y actividades."""
    villa = Sector.objects.create(comuna=comuna, nombre="Villa Las Araucarias Ficticia", tipo="villa")
    etapa = Sector.objects.create(comuna=comuna, nombre="Etapa 2 Ficticia", tipo="etapa", sector_padre=villa)
    comunidad = Comunidad.objects.create(nombre="Comunidad Los Aromos Ficticia", comuna=comuna, sector=villa)
    junta = Organizacion.objects.create(
        nombre="Junta Vecinal Ficticia N° 1", tipo=TipoOrganizacion.objects.get(codigo="junta_de_vecinos_o_comite"), direccion=_dir(comuna, villa)
    )
    Contacto.objects.create(organizacion=junta, tipo=TipoContacto.objects.get(codigo="telefono_movil"), valor="+56922223333")
    Contacto.objects.create(organizacion=junta, tipo=TipoContacto.objects.get(codigo="correo"), valor="junta.ficticia@ejemplo.invalid")
    Organizacion.objects.create(
        nombre="Club Antiguo Ficticio", tipo=TipoOrganizacion.objects.get(codigo="junta_de_vecinos_o_comite"), direccion=_dir(comuna), activo=False
    )
    vecina = crear_persona("Rosa Ficticia", apellido_paterno="Araya", rut_numero=RUT, rut_dv=calcular_dv(RUT))
    vecina.perfiles.add(PerfilPersona.objects.get(codigo="vecino_o_vecina_de_la_comunidad"))
    PersonaDireccion.objects.create(persona=vecina, direccion=_dir(comuna, etapa), tipo="residencia", vigente=True)
    ComunidadMiembro.objects.create(comunidad=comunidad, persona=vecina)
    Contacto.objects.create(persona=vecina, tipo=TipoContacto.objects.get(codigo="correo"), valor=CORREO)
    Contacto.objects.create(persona=vecina, tipo=TipoContacto.objects.get(codigo="telefono_movil"), valor=TELEFONO)
    PerfilSensible.objects.create(persona=vecina, alergias="Alergia ficticia reservada", condiciones_salud="Condición ficticia reservada", genero="mujer")
    ContactoEmergencia.objects.create(persona=vecina, nombre="Contacto Emergencia Ficticio", telefono="+56 9 0000 0000")
    pasada = crear_actividad("Taller de huertos Ficticio")
    programar(pasada, gestor, HOY - timedelta(days=10), None, None, lugar, None)
    _participa(pasada, vecina)
    SesionActividad.objects.create(actividad=pasada, numero=1, fecha=HOY - timedelta(days=10), asistentes_sin_registro=12)
    proxima = crear_actividad("Charla de reciclaje Ficticia")
    programar(proxima, gestor, HOY + timedelta(days=5), None, None, lugar, None)
    return {"comuna": comuna, "villa": villa, "etapa": etapa, "comunidad": comunidad, "junta": junta, "vecina": vecina, "lugar": lugar}


def _url_comuna(comuna):
    return reverse("territorio:comuna_exportar", args=[comuna.pk])


def _url_sector(sector):
    return reverse("territorio:sector_exportar", args=[sector.pk])


# ---------- hojas ----------

def test_exportacion_con_7_hojas(cg, mundo):
    comuna = mundo["comuna"]
    r = cg.get(_url_comuna(comuna))
    libro = _libro(r)
    assert libro.sheetnames == HOJAS + ["Aviso"]
    assert r["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert r["Content-Disposition"] == f'attachment; filename="territorio_comuna-ficticia_{HOY:%Y-%m-%d}.xlsx"'
    assert "Contiene datos personales" in _filas(libro, "Aviso")[0][0]

    resumen = dict(f[:2] for f in _filas(libro, "Resumen")[1:])
    assert _filas(libro, "Resumen")[0] == ["Indicador", "Valor"]
    assert resumen["Territorio"] == "Comuna Ficticia"
    assert (resumen["Comunidades activas"], resumen["Organizaciones activas"], resumen["Personas residentes"], resumen["Lugares activos"]) == (1, 1, 1, 1)
    assert (resumen["Actividades de los últimos 12 meses"], resumen["Actividades en total"], resumen["Total alcanzado"]) == (1, 2, 13)

    sectores = _filas(libro, "Sectores")
    assert sectores[0] == ["Sector", "Tipo", "Comunidades", "Organizaciones", "Lugares", "Personas residentes", "Actividades"]
    assert [f[0] for f in sectores[1:]] == ["Villa Las Araucarias Ficticia", "Villa Las Araucarias Ficticia › Etapa 2 Ficticia", "Sin sector"]
    assert sectores[1][1:] == ["Villa", 1, 1, 0, 0, 0] and sectores[2][1:] == ["Etapa", 0, 0, 0, 1, 0]

    comunidades = _filas(libro, "Comunidades")
    assert comunidades[0] == ["Comunidad", "Tipo", "Sector", "Organización que la representa", "Integrantes vigentes"]
    assert comunidades[1] == ["Comunidad Los Aromos Ficticia", "Vecinal o barrial", "Villa Las Araucarias Ficticia", None, 1]  # una celda vacía se lee como None

    organizaciones = _filas(libro, "Organizaciones")
    assert organizaciones[0] == ["Nombre", "Tipo", "Sector", "Correo o teléfono principal", "Activa"]
    assert organizaciones[1] == ["Junta Vecinal Ficticia N° 1", "Junta de vecinos o comité", "Villa Las Araucarias Ficticia", "junta.ficticia@ejemplo.invalid", "Sí"]
    assert organizaciones[2][0] == "Club Antiguo Ficticio" and organizaciones[2][4] == "No"  # lo desactivado sigue a la vista, al final

    personas = _filas(libro, "Personas residentes")
    assert personas[0] == ["Nombre", "Perfiles", "Comunidad", "RUT", "Correo", "Teléfono"]  # el Gestor tiene personas.exportar_persona
    assert personas[1][:3] == ["Rosa Ficticia Araya", "Vecino o vecina de la comunidad", "Comunidad Los Aromos Ficticia"]

    actividades = _filas(libro, "Actividades")
    assert actividades[0] == ["Nombre", "Estado", "Fecha", "Lugar"]
    assert [f[0] for f in actividades[1:]] == ["Charla de reciclaje Ficticia", "Taller de huertos Ficticio"]  # fecha de la agenda vigente, de la más nueva a la más antigua
    charla, taller = actividades[1], actividades[2]
    assert (charla[2].date(), charla[3]) == (HOY + timedelta(days=5), "Sala Ficticia")
    assert (taller[2].date(), taller[3]) == (HOY - timedelta(days=10), "Sala Ficticia")
    assert charla[1] and taller[1]

    alcance = _filas(libro, "Alcance")
    assert alcance[0] == ["Indicador", "Valor"]
    valores = dict(f[:2] for f in alcance[1:] if f[0] and f[1] is not None)
    assert (valores["Total alcanzado"], valores["Asistieron (personas registradas)"], valores["Asistentes sin registro"], valores["Actividades"]) == (13, 1, 12, 2)
    assert "Alergia" not in " ".join(_celdas(libro))
    listas = [f for f in alcance if f[0] in ("Instituciones de procedencia del equipo", "Comunidades de procedencia", "Participaciones por rol")]
    assert [f[1] for f in listas] == ["Participaciones"] * 3
    assert ["Sin datos", None] in alcance and ["Participante", 1] in alcance  # una lista sin filas lo dice, en vez de quedar en blanco


def test_exportacion_de_sector_incluye_a_sus_hijos(cg, mundo):
    villa = mundo["villa"]
    r = cg.get(_url_sector(villa))
    libro = _libro(r)
    assert libro.sheetnames == HOJAS + ["Aviso"]
    assert r["Content-Disposition"] == f'attachment; filename="territorio_villa-las-araucarias-ficticia_{HOY:%Y-%m-%d}.xlsx"'
    resumen = dict(f[:2] for f in _filas(libro, "Resumen")[1:])
    assert (resumen["Territorio"], resumen["Comuna"]) == ("Villa Las Araucarias Ficticia", "Comuna Ficticia")
    assert (resumen["Comunidades activas"], resumen["Organizaciones activas"], resumen["Personas residentes"]) == (1, 1, 1)  # la residente vive en la etapa
    sectores = [f[0] for f in _filas(libro, "Sectores")[1:]]
    assert sectores == ["Villa Las Araucarias Ficticia", "Villa Las Araucarias Ficticia › Etapa 2 Ficticia"]  # sin fila «Sin sector»
    assert _filas(libro, "Personas residentes")[1][0] == "Rosa Ficticia Araya"
    assert [f[0] for f in _filas(libro, "Organizaciones")[1:]] == ["Junta Vecinal Ficticia N° 1"]  # la inactiva está en la comuna, sin sector


def test_exportacion_de_comuna_o_sector_inexistente_da_404(cg):
    assert cg.get(reverse("territorio:comuna_exportar", args=[999999])).status_code == 404
    assert cg.get(reverse("territorio:sector_exportar", args=[999999])).status_code == 404


def test_exportacion_de_territorio_vacio_igual_trae_las_hojas(cg, comuna):
    libro = _libro(cg.get(_url_comuna(comuna)))
    assert libro.sheetnames == HOJAS + ["Aviso"]
    assert [f[0] for f in _filas(libro, "Sectores")[1:]] == ["Sin sector"]
    for hoja in ("Comunidades", "Organizaciones", "Personas residentes", "Actividades"):
        assert len(_filas(libro, hoja)) == 1  # solo los encabezados


# ---------- permisos ----------

def test_exportacion_403_sin_permiso(cs, cg, mundo, gestor):
    comuna, villa = mundo["comuna"], mundo["villa"]
    assert cs.get(_url_comuna(comuna)).status_code == 403
    assert cs.get(_url_sector(villa)).status_code == 403
    # ver el territorio no basta: hace falta el permiso de exportar
    lector = _cliente(_usuario("lector_territorio", *VISTAS))
    assert lector.get(_url_comuna(comuna)).status_code == 403
    assert lector.get(_url_sector(villa)).status_code == 403
    assert not Auditoria.objects.filter(accion="resumen_territorial_exportado").exists()
    # el Gestor lo tiene de fábrica y el botón aparece solo para quien puede
    assert gestor.has_perm("territorio.exportar_territorio")
    assert _url_comuna(comuna) in cg.get(comuna.get_absolute_url()).content.decode()
    assert _url_sector(villa) in cg.get(villa.get_absolute_url()).content.decode()
    assert _url_comuna(comuna) not in lector.get(comuna.get_absolute_url()).content.decode()
    assert _url_sector(villa) not in lector.get(villa.get_absolute_url()).content.decode()


def test_exportacion_exige_poder_ver_el_territorio(mundo):
    solo_exportar = _cliente(_usuario("solo_exportar", "exportar_territorio"))
    assert solo_exportar.get(_url_comuna(mundo["comuna"])).status_code == 403
    assert solo_exportar.get(_url_sector(mundo["villa"])).status_code == 403


def test_exportacion_sin_rut_ni_contactos_sin_exportar_persona(mundo):
    con = _cliente(_usuario("exporta_personas", *VISTAS, "exportar_territorio", "exportar_persona"))
    sin = _cliente(_usuario("no_exporta_personas", *VISTAS, "exportar_territorio"))

    libro = _libro(con.get(_url_comuna(mundo["comuna"])))
    personas = _filas(libro, "Personas residentes")
    assert personas[0] == ["Nombre", "Perfiles", "Comunidad", "RUT", "Correo", "Teléfono"]
    assert personas[1] == ["Rosa Ficticia Araya", "Vecino o vecina de la comunidad", "Comunidad Los Aromos Ficticia", formatear_rut(RUT, calcular_dv(RUT)), CORREO, TELEFONO]

    for url in (_url_comuna(mundo["comuna"]), _url_sector(mundo["villa"])):
        libro = _libro(sin.get(url))
        personas = _filas(libro, "Personas residentes")
        assert personas[0] == ["Nombre", "Perfiles", "Comunidad"]  # las columnas no existen, no van vacías
        assert personas[1] == ["Rosa Ficticia Araya", "Vecino o vecina de la comunidad", "Comunidad Los Aromos Ficticia"]
        todo = " ".join(_celdas(libro))
        for reservado in (formatear_rut(RUT, calcular_dv(RUT)), str(RUT), CORREO, TELEFONO):
            assert reservado not in todo, reservado


def test_exportacion_nunca_trae_datos_sensibles(cg, ca, mundo):
    for cliente in (cg, ca):
        libro = _libro(cliente.get(_url_comuna(mundo["comuna"])))
        todo = " ".join(_celdas(libro))
        for sensible in ("Alergia ficticia reservada", "Condición ficticia reservada", "Contacto Emergencia Ficticio", "+56 9 0000 0000", "Mujer", "mujer"):
            assert sensible not in todo, sensible
        assert [h for h in libro.sheetnames if "ensible" in h or "mergencia" in h or "erificaci" in h] == []


def test_exportacion_sin_ver_personas_omite_la_hoja(mundo):
    sin_personas = _cliente(_usuario("sin_ver_personas", *[v for v in VISTAS if v != "view_persona"], "exportar_territorio", "exportar_persona"))
    for url in (_url_comuna(mundo["comuna"]), _url_sector(mundo["villa"])):
        libro = _libro(sin_personas.get(url))
        assert libro.sheetnames == [h for h in HOJAS if h != "Personas residentes"] + ["Aviso"]  # 6 hojas y el aviso
        todo = " ".join(_celdas(libro))
        assert "Rosa Ficticia" not in todo and CORREO not in todo
        assert dict(f[:2] for f in _filas(libro, "Resumen")[1:])["Personas residentes"] == 1  # el conteo no identifica a nadie


@pytest.mark.parametrize("sin_permiso, hoja", [
    ("view_comunidad", "Comunidades"), ("view_organizacion", "Organizaciones"), ("view_actividad", "Actividades"), ("view_participacion", "Alcance"),
])
def test_exportacion_omite_las_hojas_que_el_usuario_no_veria_en_pantalla(mundo, sin_permiso, hoja):
    usuario = _cliente(_usuario("recortado", *[v for v in VISTAS if v != sin_permiso], "exportar_territorio"))
    libro = _libro(usuario.get(_url_comuna(mundo["comuna"])))
    assert libro.sheetnames == [h for h in HOJAS if h != hoja] + ["Aviso"]
    if hoja == "Alcance":
        assert "Total alcanzado" not in [f[0] for f in _filas(libro, "Resumen")]


def test_exportacion_neutraliza_formulas_en_las_celdas(cg, mundo):
    Comunidad.objects.create(nombre='=HYPERLINK("http://ejemplo.invalid","clic")', comuna=mundo["comuna"])
    Organizacion.objects.create(nombre="+cmd|' /C calc'!A0", tipo=TipoOrganizacion.objects.get(codigo="empresa_o_partner"), direccion=_dir(mundo["comuna"]))
    libro = _libro(cg.get(_url_comuna(mundo["comuna"])))
    textos = [str(v) for h in libro for f in h.iter_rows(values_only=True) for v in f if isinstance(v, str)]
    assert not [t for t in textos if t[:1] in "=@" or (t[:1] == "+" and not t[1:].isdigit())]  # un teléfono "+569…" es un dato, no una fórmula
    assert any(t.startswith("'=HYPERLINK") for t in textos) and any(t.startswith("'+cmd") for t in textos)


# ---------- auditoría ----------

def test_exportacion_queda_en_auditoria(cg, gestor, mundo):
    cg.get(_url_comuna(mundo["comuna"]))
    cg.get(_url_sector(mundo["villa"]))
    comuna, sector = Auditoria.objects.filter(accion="resumen_territorial_exportado").order_by("id")
    assert (comuna.usuario, comuna.usuario_nombre, comuna.objeto) == (gestor, gestor.username, "Comuna Ficticia")
    assert sector.objeto == "Villa Las Araucarias Ficticia" and sector.usuario == gestor
    assert "Personas residentes" in comuna.detalle


def test_auditoria_anota_si_salieron_rut_y_contactos(mundo):
    sin = _cliente(_usuario("no_exporta_personas", *VISTAS, "exportar_territorio"))
    con = _cliente(_usuario("exporta_personas", *VISTAS, "exportar_territorio", "exportar_persona"))
    sin.get(_url_comuna(mundo["comuna"]))
    con.get(_url_comuna(mundo["comuna"]))
    sin_rut, con_rut = Auditoria.objects.filter(accion="resumen_territorial_exportado").order_by("id")
    assert "con RUT y contactos" not in sin_rut.detalle and "con RUT y contactos" in con_rut.detalle


# ---------- rendimiento ----------

def test_exportacion_usa_un_numero_fijo_de_consultas(cg, mundo, crear_persona, crear_actividad, gestor):
    comuna, villa, comunidad, lugar = mundo["comuna"], mundo["villa"], mundo["comunidad"], mundo["lugar"]
    exporta = _cliente(_usuario("exporta_personas", *VISTAS, "exportar_territorio", "exportar_persona"))

    def consultas(cliente, url):
        with CaptureQueriesContext(connection) as ctx:
            assert cliente.get(url).status_code == 200
        return len(ctx)

    base = {(c, u): consultas(c, u) for c in (cg, exporta) for u in (_url_comuna(comuna), _url_sector(villa))}
    perfil = PerfilPersona.objects.get(codigo="vecino_o_vecina_de_la_comunidad")
    for i in range(30):
        p = crear_persona(f"Vecina{i}", rut_numero=20000000 + i, rut_dv=calcular_dv(20000000 + i))
        p.perfiles.add(perfil)
        PersonaDireccion.objects.create(persona=p, direccion=_dir(comuna, villa), tipo="residencia", vigente=True)
        ComunidadMiembro.objects.create(comunidad=Comunidad.objects.create(nombre=f"Comunidad {i}", comuna=comuna, sector=villa), persona=p)
        Contacto.objects.create(persona=p, tipo=TipoContacto.objects.get(codigo="correo"), valor=f"vecina{i}@ejemplo.invalid")
        Contacto.objects.create(persona=p, tipo=TipoContacto.objects.get(codigo="telefono_movil"), valor=f"+5693000{i:04d}")
        org = Organizacion.objects.create(nombre=f"Organización {i}", tipo=TipoOrganizacion.objects.get(codigo="empresa_o_partner"), direccion=_dir(comuna, villa))
        Contacto.objects.create(organizacion=org, tipo=TipoContacto.objects.get(codigo="correo"), valor=f"org{i}@ejemplo.invalid")
        a = crear_actividad(f"Actividad {i}")
        programar(a, gestor, HOY - timedelta(days=i), None, None, lugar, None)
        _participa(a, p)
    for (cliente, url), antes in base.items():
        assert consultas(cliente, url) == antes, url
        assert antes <= 45
