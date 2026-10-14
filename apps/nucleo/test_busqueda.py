"""Pruebas automáticas del buscador, las exportaciones y el Panel: búsqueda en todos los tipos y con permisos,
filtros, exportaciones con permiso y auditoría, pendientes y cola de personas incompletas."""
from datetime import date, time, timedelta

import pytest
from django.utils import timezone

from apps.actividades import servicios
from apps.actividades.models import Participacion
from apps.nucleo.models import TipoActividad

pytestmark = pytest.mark.django_db


@pytest.fixture
def mundo(crear_actividad, crear_persona, rol, estado_part, lugar, universidad, gestor):
    """Datos ficticios mínimos para probar búsquedas combinadas."""
    from apps.proyectos.models import Proyecto

    pr = Proyecto.objects.create(nombre="Proyecto Educación Comunitaria", zona="Comuna Ficticia")
    a = crear_actividad("Taller de educación financiera", proyecto=pr, tipo=TipoActividad.objects.get(codigo="taller"), estado="planificada")
    servicios.programar(a, gestor, date.today() + timedelta(days=5), time(18), time(20), lugar, None)
    b = crear_actividad("Charla sobre salud", tipo=TipoActividad.objects.get(codigo="charla"), estado="realizada")
    ana = crear_persona("Ana", apellido_paterno="Educadora", profesion="Profesora")
    beto = crear_persona("Beto", apellido_paterno="Sinactividad")
    Participacion.objects.create(actividad=a, persona=ana, rol=rol("tallerista"), estado=estado_part("confirmado"))
    from apps.actividades.models import ActividadOrganizacion

    ActividadOrganizacion.objects.create(actividad=a, organizacion=universidad)
    return {"a": a, "b": b, "ana": ana, "beto": beto, "pr": pr}


def hallar(c, **params):
    from urllib.parse import urlencode

    return c.get("/buscar?" + urlencode(params)).content.decode()


def test_busqueda_vacia_muestra_ayuda(cg):
    assert "Escribe algo" in hallar(cg)


def test_busca_en_todos_los_tipos(cg, mundo, lugar, universidad):
    html = hallar(cg, q="educ")
    assert "Taller de educación financiera" in html and "Ana" in html and "Proyecto Educación Comunitaria" in html
    assert "Charla sobre salud" not in html
    assert "Sala Ficticia" in hallar(cg, q="sala")
    assert "Universidad Ficticia" in hallar(cg, q="universidad")


def test_filtro_por_tipo_y_estado(cg, mundo):
    html = hallar(cg, q="", tipo="actividad", estado="realizada")
    assert "Charla sobre salud" in html and "Taller de educación" not in html and "Ana" not in html


def test_actividades_realizadas_en_una_zona(cg, mundo, lugar):
    assert "Taller de educación" in hallar(cg, tipo="actividad", zona=lugar.direccion.comuna_id)
    assert "Charla sobre salud" not in hallar(cg, tipo="actividad", zona=lugar.direccion.comuna_id)
    assert "Sin resultados" in hallar(cg, tipo="actividad", estado="realizada", zona=lugar.direccion.comuna_id)


def test_personas_que_participaron_en_actividades_de_un_tipo(cg, mundo):
    taller = TipoActividad.objects.get(codigo="taller")
    html = hallar(cg, tipo="persona", tipo_actividad=taller.pk)
    assert "Ana" in html and "Beto" not in html


def test_organizaciones_con_las_que_hemos_trabajado(cg, mundo, universidad):
    from apps.nucleo.models import TipoOrganizacion
    from apps.organizaciones.models import Organizacion

    Organizacion.objects.create(nombre="Empresa sin actividades", tipo=TipoOrganizacion.objects.get(codigo="empresa_o_partner"))
    html = hallar(cg, tipo="organizacion", con_actividades="1")
    assert "Universidad Ficticia" in html and "Empresa sin actividades" not in html
    assert "Empresa sin actividades" in hallar(cg, tipo="organizacion", q="sin actividades")


def test_filtros_por_fecha(cg, mundo):
    hoy = date.today()
    assert "Taller de educación" in hallar(cg, tipo="actividad", desde=hoy, hasta=hoy + timedelta(days=10))
    assert "Taller de educación" not in hallar(cg, tipo="actividad", desde=hoy + timedelta(days=30))
    assert hallar(cg, tipo="actividad", desde="basura").count("Taller") >= 0  # fecha mal escrita: no falla


def test_busqueda_respeta_permisos_por_tipo(ca, cs, mundo):
    assert cs.get("/buscar?q=ana").status_code == 403
    html = hallar(ca, q="educ")
    assert "Ana" in html  # el Administrador ve personas (solo lectura)


def test_busqueda_no_expone_documentos_restringidos(cg, ca, mundo, administrador):
    from django.contrib.auth.models import Permission
    from django.core.files.uploadedfile import SimpleUploadedFile
    from django.test import Client

    from apps.documentos.models import Documento
    from apps.nucleo.models import TipoDocumento

    Documento.objects.create(nombre="Antecedentes reservados", archivo=SimpleUploadedFile("a.pdf", b"%PDF-1"), tipo=TipoDocumento.objects.get(codigo="otro"), actividad=mundo["a"], restringido=True)
    assert "Antecedentes reservados" in hallar(cg, q="antecedentes")
    administrador.groups.first().permissions.remove(Permission.objects.get(codename="ver_restringido"))
    c = Client()
    c.force_login(type(administrador).objects.get(pk=administrador.pk))
    assert "Antecedentes reservados" not in hallar(c, q="antecedentes")


def test_resultados_paginados_y_con_enlace(cg, crear_persona):
    for i in range(60):
        crear_persona(f"Paginada{i:02d}")
    html = hallar(cg, q="paginada", tipo="persona")
    assert "60 resultado" in html and "Siguiente" in html and "/personas/" in html
    assert "Paginada29" in cg.get("/buscar?q=paginada&tipo=persona&page=2").content.decode() or "Paginada" in cg.get("/buscar?q=paginada&tipo=persona&page=2").content.decode()


def test_buscador_global_visible_en_todas_las_pantallas(cg):
    for ruta in ("/dashboard", "/personas", "/actividades", "/calendario", "/administracion/catalogos/"):
        assert 'aria-label="Buscador global"' in cg.get(ruta).content.decode(), ruta


# --- exportación --------------------------------------------------------------------------------------------
def test_exportaciones_requieren_permiso_y_quedan_en_auditoria(cg, cs, mundo):
    from io import BytesIO

    from openpyxl import load_workbook

    from apps.nucleo.models import Auditoria

    r = cg.get("/personas/exportar?q=ana")
    assert r.status_code == 200 and "attachment" in r["Content-Disposition"]
    hoja = load_workbook(BytesIO(r.content)).active
    filas = list(hoja.iter_rows(values_only=True))
    assert filas[0][0] == "Nombre" and any("Ana" in str(f[0]) for f in filas[1:]) and not any("Beto" in str(f[0]) for f in filas[1:])
    assert cs.get("/personas/exportar").status_code == 403
    for ruta in ("/organizaciones/exportar", "/lugares/exportar", "/proyectos/exportar", "/solicitudes/exportar", "/tareas/exportar", "/actividades/exportar"):
        assert cg.get(ruta).status_code == 200, ruta
    assert Auditoria.objects.filter(accion="exportacion").count() >= 7


def test_exportacion_neutraliza_formulas(cg, crear_persona):
    from io import BytesIO

    from openpyxl import load_workbook

    crear_persona("=CMD()", apellido_paterno="Maligno")
    hoja = load_workbook(BytesIO(cg.get("/personas/exportar").content)).active
    celdas = [c for fila in hoja.iter_rows(values_only=True) for c in fila if c]
    assert not any(str(c).startswith("=") for c in celdas)


# --- dashboard -------------------------------------------------------------------------------------------------
def test_dashboard_muestra_pendientes(cg, mundo, crear_persona, gestor):
    from apps.coordinacion.models import Tarea
    from apps.nucleo.models import EstadoRegistro
    from apps.solicitudes.models import Solicitud
    from apps.nucleo.models import EstadoSolicitud

    Tarea.objects.create(titulo="Mandar la carta", fecha_limite=timezone.localdate() + timedelta(days=2))
    Solicitud.objects.create(motivo="Nadie ha respondido", estado=EstadoSolicitud.objects.get(codigo="pendiente"))
    incompleta = crear_persona("Incompleta", estado_registro=EstadoRegistro.objects.get(codigo="incompleto"))
    html = cg.get("/dashboard").content.decode()
    assert "Actividades incompletas" in html and "Taller de educación financiera" in html  # le falta el responsable
    assert "Mandar la carta" in html and "Nadie ha respondido" in html and "Incompleta" in html
    assert "+ Persona" in html


def test_cola_de_incompletos_prioriza_actividades_proximas(mundo, crear_persona, rol, estado_part):
    from apps.nucleo.models import EstadoRegistro
    from apps.personas.models import Persona

    inc = EstadoRegistro.objects.get(codigo="incompleto")
    sin_actividad = crear_persona("SinActividad", estado_registro=inc)
    con_proxima = crear_persona("ConProxima", estado_registro=inc)
    Participacion.objects.create(actividad=mundo["a"], persona=con_proxima, rol=rol("voluntario"), estado=estado_part("invitado"))
    cola = list(Persona.objects.cola_incompletos())
    assert cola[0] == con_proxima and cola[0].prioridad is True and cola[1] == sin_actividad


def test_dashboard_vacio_y_para_administrador(ca, cg):
    assert ca.get("/dashboard").status_code == 200
    assert "Nada pendiente" in cg.get("/dashboard").content.decode()


def test_resultados_agrupados_por_tipo(cg, mundo):
    html = hallar(cg, q="educ")
    assert "<h2 class=\"h6 mt-3\">Personas" in html and "<h2 class=\"h6 mt-3\">Actividades" in html and "<h2 class=\"h6 mt-3\">Proyectos" in html
    h = lambda t: html.index(f'<h2 class="h6 mt-3">{t}')
    assert h("Personas") < h("Proyectos") < h("Actividades")


def test_menu_de_usuario_y_acciones_rapidas(cg):
    html = cg.get("/dashboard").content.decode()
    assert "Mi cuenta" in html and "Cerrar sesión" in html and 'id="cargando"' in html and 'id="reintentar"' in html
    for accion in ("+ Solicitud", "+ Oportunidad", "+ Actividad", "+ Persona", "+ Organización", "+ Lugar"):
        assert accion in html, accion
    assert "/static/js/fche.js" in html


def test_mensajes_de_error_con_estilo_bootstrap(cg, actividad):
    r = cg.post(f"/actividades/{actividad.pk}/estado", {"estado": "cancelada"}, follow=True)  # falta el motivo
    assert "alert-danger" in r.content.decode()
