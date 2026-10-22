"""F22 — Ingesta ampliada: personas, estudiantes, comunidades, público de la actividad y bitácora de sesiones. Datos ficticios."""
from datetime import date, time, timedelta

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone

from apps.actividades.models import ActividadComunidad, Participacion, SesionActividad, SesionAsistencia
from apps.comunidades.models import Comunidad, ComunidadMiembro
from apps.nucleo.models import AreaTematica, PerfilPersona, RolParticipacion
from apps.personas.models import Apoderado, Consentimiento, PerfilSensible, Persona, PersonaFormacion

pytestmark = pytest.mark.django_db
HOY = timezone.localdate()


def nacido_hace(anios):
    return HOY.replace(year=HOY.year - anios) - timedelta(days=1)


@pytest.fixture
def comunidad(comuna):
    return Comunidad.objects.create(nombre="Vecinos del Sector Ficticio", comuna=comuna, tipo="vecinal")


# --- 1. Persona: ficha completa por secciones ---------------------------------------------------
def test_ficha_completa_guarda_identificacion_perfil_e_intereses(cg):
    perfil = PerfilPersona.objects.get(codigo="estudiante")
    area = AreaTematica.objects.get(codigo="medioambiente")
    r = cg.post("/personas/nueva", {
        "modo": "completa", "nombres": "Lía", "apellido_paterno": "Ficticia", "fecha_nacimiento": "2004-05-10", "nacionalidad": "Chilena",
        "nombre_social": "Lía", "nivel_educacional": "universitaria_incompleta", "situacion_ocupacional": "estudia", "perfiles": [perfil.pk],
        "areas_interes": [area.pk], "habilidades": "Fotografía y huertos urbanos", "idiomas": "Inglés básico", "correo": "lia@ejemplo.invalid",
    })
    assert r.status_code == 302
    p = Persona.objects.get(nombres="Lía")
    assert p.fecha_nacimiento == date(2004, 5, 10) and p.nacionalidad == "Chilena" and p.edad >= 21 and not p.es_menor
    assert list(p.perfiles.all()) == [perfil] and list(p.areas_interes.all()) == [area]
    html = cg.get(p.get_absolute_url()).content.decode()
    assert f"{p.edad} años" in html and "Estudiante" in html


def test_formulario_por_secciones_y_fecha_futura_rechazada(cg, persona):
    html = cg.get(f"/personas/{persona.pk}/editar").content.decode()
    for seccion in ("Identificación", "Perfil y ocupación", "Intereses y habilidades"):
        assert f'<legend class="fche-bloque__titulo">{seccion}</legend>' in html
    r = cg.post(f"/personas/{persona.pk}/editar", {"nombres": persona.nombres, "fecha_nacimiento": (HOY + timedelta(days=3)).isoformat()})
    assert r.status_code == 200 and "no puede ser futura" in r.content.decode()


def test_edad_y_menor_de_edad(crear_persona):
    assert crear_persona(fecha_nacimiento=nacido_hace(17)).es_menor
    assert not crear_persona(fecha_nacimiento=nacido_hace(18)).es_menor
    assert crear_persona().edad is None


# --- 2. Datos sensibles -------------------------------------------------------------------------
def test_datos_sensibles_exigen_consentimiento(cg, persona):
    url = f"/personas/{persona.pk}/sensibles/editar"
    r = cg.post(url, {"pueblo_originario": "mapuche", "alergias": "Maní"})
    assert r.status_code == 400 and "consentimiento" in r.content.decode() and not PerfilSensible.objects.exists()
    r = cg.post(url, {"pueblo_originario": "mapuche", "alergias": "Maní", "consentimiento": "on", "consentimiento_fecha": HOY.isoformat(), "consentimiento_medio": "papel"})
    assert r.status_code == 302
    assert persona.perfil_sensible.alergias == "Maní" and persona.consentimiento_vigente("sensibles")
    # Con consentimiento vigente ya no se vuelve a pedir
    assert "consentimiento_fecha" not in cg.get(url).content.decode()


def test_datos_sensibles_permisos_y_ajuste(cg, cs, persona, settings):
    assert cs.get(f"/personas/{persona.pk}/sensibles/editar").status_code == 403
    # Los datos sensibles son una sección de la pestaña Privacidad 🔒: se mira esa pestaña, no la barra de pestañas.
    privacidad = cg.get(f"{persona.get_absolute_url()}?tab=privacidad").content.decode().split('id="pestana"')[1]
    assert "Datos sensibles (Ley 21.719)" in privacidad  # el aviso propio de la sección (el consentimiento «Datos sensibles» sigue aparte)
    settings.DATOS_SENSIBLES = False
    assert cg.get(f"/personas/{persona.pk}/sensibles/editar").status_code == 404
    privacidad = cg.get(f"{persona.get_absolute_url()}?tab=privacidad").content.decode().split('id="pestana"')[1]
    assert "Datos sensibles (Ley 21.719)" not in privacidad


def test_gestor_sin_permiso_de_sensibles_no_ve_la_pestana(gestor, persona):
    from django.contrib.auth.models import Permission
    from django.test import Client

    gestor.groups.first().permissions.remove(Permission.objects.get(codename="ver_sensibles"))
    c = Client()
    c.force_login(gestor)
    tabs = c.get(persona.get_absolute_url()).content.decode().split("nav-tabs")[1].split("</ul>")[0]
    assert "Datos sensibles" not in tabs and "Apoderados" not in tabs
    assert c.get(f"/personas/{persona.pk}/sensibles/editar").status_code == 403


# --- 3. Consentimientos y menores ---------------------------------------------------------------
def test_menor_sin_apoderado_ni_autorizacion_muestra_avisos(cg, crear_persona):
    menor = crear_persona("Tomás", fecha_nacimiento=nacido_hace(12))
    html = cg.get(menor.get_absolute_url()).content.decode()
    assert "Menor de edad" in html and "no tiene apoderado registrado" in html and "falta la autorización del apoderado" in html
    madre = crear_persona("Rosa")
    Apoderado.objects.create(persona=menor, apoderado=madre, parentesco="madre", tutor_legal=True)
    Consentimiento.objects.create(persona=menor, tipo="apoderado", otorgado_por=madre)
    Consentimiento.objects.create(persona=menor, tipo="datos", otorgado_por=madre)
    assert menor.avisos_consentimiento() == []


def test_registrar_y_revocar_consentimiento(cg, gestor, persona):
    r = cg.post(f"/personas/{persona.pk}/consentimientos/nuevo", {"tipo": "imagen", "otorgado": "on", "fecha": HOY.isoformat(), "medio": "papel"})
    assert r.status_code == 302
    c = Consentimiento.objects.get()
    assert c.registrado_por == gestor and c.vigente and persona.consentimiento_vigente("imagen")
    cg.post(f"/personas/{persona.pk}/consentimientos/{c.pk}/editar", {"tipo": "imagen", "otorgado": "on", "fecha": HOY.isoformat(), "medio": "papel", "revocado_en": HOY.isoformat()})
    assert not persona.consentimiento_vigente("imagen") and Consentimiento.objects.count() == 1  # se revoca, no se borra
    html = cg.get(f"{persona.get_absolute_url()}?tab=consentimientos").content.decode()
    assert "Revocado" in html


# --- 4. Estudiante --------------------------------------------------------------------------------
def test_formacion_ampliada_del_estudiante(cg, persona, universidad, crear_persona):
    docente = crear_persona("Profesora")
    r = cg.post(f"/personas/{persona.pk}/formacion/nuevo", {
        "organizacion": universidad.pk, "estado": "estudiante", "sede": "Campus Ficticio", "jornada": "vespertina", "nivel_actual": 3,
        "asignatura": "Aprendizaje + servicio", "docente": docente.pk, "horas_comprometidas": 40, "vigente": "on",
    })
    assert r.status_code == 302
    f = PersonaFormacion.objects.get()
    assert f.sede == "Campus Ficticio" and f.docente == docente and f.horas_comprometidas == 40
    assert "Campus Ficticio" in cg.get(f"{persona.get_absolute_url()}?tab=formacion").content.decode()


# --- 5. Comunidades -----------------------------------------------------------------------------
def test_crear_comunidad_con_padron_y_verla_en_la_comuna(cg, comuna, persona):
    r = cg.post("/comunidades/nueva", {"nombre": "Club de adulto mayor Ficticio", "tipo": "personas_mayores", "comuna": comuna.pk, "poblacion_estimada": 40})
    assert r.status_code == 302
    c = Comunidad.objects.get()
    assert cg.post(f"/comunidades/{c.pk}/integrantes/nuevo", {"persona": persona.pk, "rol": "dirigente", "vigente": "on"}).status_code == 302
    assert cg.post(f"/comunidades/{c.pk}/integrantes/nuevo", {"persona": persona.pk, "rol": "integrante", "vigente": "on"}).status_code == 200  # ya integra
    assert ComunidadMiembro.objects.count() == 1
    html = cg.get(f"{c.get_absolute_url()}?tab=integrantes").content.decode()
    assert "Ana María Pérez Soto" in html and "Dirigente/a" in html
    assert "Club de adulto mayor Ficticio" in cg.get(f"{comuna.get_absolute_url()}?tab=comunidades").content.decode()
    assert "Club de adulto mayor Ficticio" in cg.get(f"{persona.get_absolute_url()}?tab=comunidades").content.decode()
    assert "Club de adulto mayor Ficticio" in cg.get("/comunidades").content.decode()


def test_sector_de_otra_comuna_rechazado(cg, comuna):
    from apps.territorio.models import Comuna, Sector

    otra = Comuna.objects.create(provincia=comuna.provincia, codigo="99002", nombre="Otra Comuna")
    s = Sector.objects.create(comuna=otra, nombre="Villa Ajena")
    r = cg.post("/comunidades/nueva", {"nombre": "Mal ubicada", "tipo": "vecinal", "comuna": comuna.pk, "sector": s.pk})
    assert r.status_code == 200 and not Comunidad.objects.exists()


def test_comunidades_y_permisos(cs, ca, comunidad):
    assert cs.get("/comunidades").status_code == 403
    assert ca.get(comunidad.get_absolute_url()).status_code == 200
    assert ca.get("/comunidades/nueva").status_code == 403  # el Administrador ve, no opera


# --- 6. Actividad: público, comunidades y procedencia -----------------------------------------
def test_actividad_para_una_comunidad_y_procedencia_automatica(cg, comunidad, persona, universidad, rol, estado_part, crear_persona):
    r = cg.post(f"/actividades/nueva?comunidad={comunidad.pk}", {"nombre": "Taller de huertos", "publico_objetivo": "Vecinos del sector", "beneficiarios_estimados": 30})
    assert r.status_code == 302
    a = ActividadComunidad.objects.get().actividad
    assert a.publico_objetivo == "Vecinos del sector" and a.beneficiarios_estimados == 30
    vecina = crear_persona("Vecina")
    ComunidadMiembro.objects.create(comunidad=comunidad, persona=vecina)
    PersonaFormacion.objects.create(persona=persona, organizacion=universidad)
    p1 = Participacion.objects.create(actividad=a, persona=persona, rol=rol("estudiante"), estado=estado_part("confirmado"))
    p2 = Participacion.objects.create(actividad=a, persona=vecina, rol=rol("participante"), estado=estado_part("confirmado"))
    assert p1.institucion == universidad and p2.comunidad == comunidad
    html = cg.get(f"{a.get_absolute_url()}?tab=participantes").content.decode()
    assert "Equipo ejecutor" in html and "Asistentes" in html and "Universidad Ficticia" in html
    assert "Vecinos del Sector Ficticio" in cg.get(f"{a.get_absolute_url()}?tab=publico").content.decode()


def test_grupo_de_los_roles():
    assert RolParticipacion.objects.get(codigo="tallerista").grupo == "ejecutor"
    assert RolParticipacion.objects.get(codigo="beneficiario").grupo == "asistente"
    assert RolParticipacion.objects.get(codigo="contacto_institucional").grupo == "contacto"


# --- 7. Bitácora de sesiones --------------------------------------------------------------------
def test_sesiones_visibles_sin_segunda_fase_con_minuta_y_asistencia(cg, sin_segunda_fase, actividad, persona, rol, estado_part):
    responsable = persona
    actividad.responsable = responsable
    actividad.save()
    part = Participacion.objects.create(actividad=actividad, persona=persona, rol=rol("tallerista"), estado=estado_part("confirmado"))
    r = cg.post(f"/actividades/{actividad.pk}/sesiones/nuevo", {
        "numero": 1, "nombre": "Diagnóstico participativo", "fecha": HOY.isoformat(), "hora_inicio": "10:00", "hora_fin": "12:00",
        "facilitadores": [persona.pk], "desarrollo": "Se mapearon los espacios del barrio.", "avances": "Priorizaron la plaza.", "acuerdos": "Limpieza el sábado.",
        "asistentes_sin_registro": 12, "conteo_mujeres": 7, "conteo_hombres": 5, "vigente": "on",
    })
    assert r.status_code == 302
    s = SesionActividad.objects.get()
    assert s.redactada_por.username == "gestor_prueba" and [f.persona for f in s.facilitadores.all()] == [persona] and s.duracion_horas == 2
    cg.post(f"/actividades/{actividad.pk}/sesiones/{s.pk}/asistencia", {f"asistio_{part.pk}": "si"})
    assert SesionAsistencia.objects.get().asistio is True and s.total_asistentes == 13
    html = cg.get(s.get_absolute_url()).content.decode()
    assert "Priorizaron la plaza." in html and "Limpieza el sábado." in html and "Asistentes sin registro individual: 12" in html
    tab = cg.get(f"{actividad.get_absolute_url()}?tab=sesiones").content.decode()
    assert "Diagnóstico participativo" in tab and "13 asistentes" in tab


def test_conteos_no_pueden_superar_el_total(cg, actividad):
    r = cg.post(f"/actividades/{actividad.pk}/sesiones/nuevo", {"numero": 1, "asistentes_sin_registro": 5, "conteo_menores": 4, "conteo_adultos": 3, "vigente": "on"})
    assert r.status_code == 200 and "supera el total" in r.content.decode() and not SesionActividad.objects.exists()


def test_evidencia_de_la_sesion(cg, actividad):
    from apps.documentos.models import Documento
    from apps.nucleo.models import TipoDocumento

    s = SesionActividad.objects.create(actividad=actividad, numero=1)
    r = cg.post("/documentos/nuevo", {"nombre": "Foto del taller", "archivo": SimpleUploadedFile("foto.pdf", b"%PDF-1.4 ficticio"), "tipo": TipoDocumento.objects.first().pk,
                                     "fecha": HOY.isoformat(), "sesion": s.pk})
    assert r.status_code == 302 and r["Location"] == s.get_absolute_url()
    d = Documento.objects.get()
    assert d.sesion == s and d.actividad == actividad
    assert "Foto del taller" in cg.get(s.get_absolute_url()).content.decode()


def test_minutas_en_el_buscador(cg, actividad):
    SesionActividad.objects.create(actividad=actividad, numero=1, acuerdos="Instalar composteras comunitarias", fecha=HOY)
    html = cg.get("/buscar?q=composteras").content.decode()
    assert '<h2 class="h6 mt-3">Minutas' in html and "Sesión 1 · Taller de prueba" in html


# --- 8. Alcance ---------------------------------------------------------------------------------
def test_alcance_cuenta_equipo_asistentes_horas_y_sin_registro(cg, actividad, lugar, gestor, crear_persona, rol, estado_part, universidad, comunidad):
    from apps.actividades import servicios
    from apps.actividades.alcance import alcance, horas_persona

    servicios.programar(actividad, gestor, HOY, time(9), time(12), lugar, None)
    ActividadComunidad.objects.create(actividad=actividad, comunidad=comunidad)
    est = crear_persona("Estudiante", fecha_nacimiento=nacido_hace(21))
    PersonaFormacion.objects.create(persona=est, organizacion=universidad)
    vecino = crear_persona("Vecino", fecha_nacimiento=nacido_hace(70))
    ComunidadMiembro.objects.create(comunidad=comunidad, persona=vecino)
    Participacion.objects.create(actividad=actividad, persona=est, rol=rol("estudiante"), estado=estado_part("asistio"), asistencia=True)
    Participacion.objects.create(actividad=actividad, persona=vecino, rol=rol("beneficiario"), estado=estado_part("asistio"), asistencia=True)
    SesionActividad.objects.create(actividad=actividad, numero=1, asistentes_sin_registro=10, conteo_mayores=4)
    a = alcance([actividad])
    assert a["equipo"] == 1 and a["asistentes"] == 1 and a["asistieron"] == 2 and a["sin_registro"] == 10 and a["total_alcanzado"] == 12
    assert a["horas"] == 6 and a["horas_equipo"] == 3 and a["instituciones"] == [(universidad, 1)] and a["comunidades"] == [(comunidad, 1)]
    assert dict(a["edades_registradas"]) == {"0 a 17": 0, "18 a 29": 1, "30 a 59": 0, "60 o más": 1}
    assert horas_persona(est) == 3
    assert "Total alcanzado" in cg.get(f"{actividad.get_absolute_url()}?tab=alcance").content.decode()
    assert "Universidad Ficticia" in cg.get(f"{comunidad.get_absolute_url()}?tab=alcance").content.decode()


def test_proyecto_con_comunidades_y_alcance(cg, comunidad, crear_actividad):
    from apps.proyectos.models import Proyecto, ProyectoComunidad

    pr = Proyecto.objects.create(nombre="Barrio Verde Ficticio")
    crear_actividad("Jornada de limpieza", proyecto=pr)
    assert cg.post(f"/proyectos/{pr.pk}/comunidades/nuevo", {"comunidad": comunidad.pk, "vigente": "on"}).status_code == 302
    assert ProyectoComunidad.objects.filter(proyecto=pr, comunidad=comunidad).exists()
    assert "Vecinos del Sector Ficticio" in cg.get(f"{pr.get_absolute_url()}?tab=comunidades").content.decode()
    assert cg.get(f"{pr.get_absolute_url()}?tab=alcance").status_code == 200
    assert "Barrio Verde Ficticio" in cg.get(f"{comunidad.get_absolute_url()}?tab=actividades").content.decode()


# --- 9. Planillas -------------------------------------------------------------------------------
def test_planilla_con_columnas_nuevas(cg, comuna, universidad):
    from apps.importaciones import procesar
    from apps.importaciones.models import Lote

    texto = (
        "Nombre;Apellido paterno;Fecha de nacimiento;Nacionalidad;Comuna;Dirección;Institución;Sede;Jornada;Año que cursa;Asignatura\n"
        f"Inés;Ficticia;03/08/2003;Peruana;{comuna.nombre};Pasaje Inventado;Universidad Ficticia;Campus Norte;Vespertina;2.º;Vinculación con el medio\n"
        "Malo;Fecha;31-02-2001;;;;;;;;\n"
    )
    cg.post("/importaciones/nueva", {"tipo": "personas", "texto": texto})
    lote = Lote.objects.get()
    assert lote.filas[1]["estado"] == "error" and "Fecha de nacimiento no reconocida" in lote.filas[1]["errores"][0]
    cg.post(f"/importaciones/{lote.pk}/revisar", {"accion": "confirmar", "d3": "omitir"})
    p = Persona.objects.get(nombres="Inés")
    assert p.fecha_nacimiento == date(2003, 8, 3) and p.nacionalidad == "Peruana"
    assert p.direcciones.get().direccion.comuna == comuna
    f = p.formaciones.get()
    assert f.sede == "Campus Norte" and f.jornada == "vespertina" and f.nivel_actual == 2 and f.asignatura == "Vinculación con el medio"
    lote.refresh_from_db()
    procesar.revertir(lote, None)
    assert not Persona.objects.filter(nombres="Inés").exists()


def test_fecha_planilla():
    from apps.importaciones.procesar import fecha_planilla

    assert fecha_planilla("03-08-2003") == date(2003, 8, 3) and fecha_planilla("2003-08-03 00:00:00") == date(2003, 8, 3)
    assert fecha_planilla("3/8/03") == date(2003, 8, 3) and fecha_planilla("31-02-2001") is None and fecha_planilla("ayer") is None


# --- 10. Banco de personas: filtros, exportación y fusión ---------------------------------------
def test_filtros_de_personas_por_perfil_comunidad_y_edad(cg, crear_persona, comunidad):
    joven = crear_persona("Joven", fecha_nacimiento=nacido_hace(20))
    mayor = crear_persona("Mayor", fecha_nacimiento=nacido_hace(65))
    joven.perfiles.add(PerfilPersona.objects.get(codigo="voluntario_o_voluntaria"))
    ComunidadMiembro.objects.create(comunidad=comunidad, persona=mayor)
    html = cg.get("/personas?edad=18_29").content.decode()
    assert "Joven" in html and "Mayor" not in html
    html = cg.get(f"/personas?perfil={PerfilPersona.objects.get(codigo='voluntario_o_voluntaria').pk}").content.decode()
    assert "Joven" in html and "Mayor" not in html
    html = cg.get(f"/personas?comunidad={comunidad.pk}").content.decode()
    assert "Mayor" in html and "Joven" not in html
    assert "Mayor" in cg.get("/personas?edad=60").content.decode()


def test_fusion_suma_perfiles_y_completa_fecha(ca, crear_persona):
    from apps.nucleo.fusion import fusionar

    a = crear_persona("Duplicada", fecha_nacimiento=date(1990, 1, 1))
    b = crear_persona("Original")
    a.perfiles.add(PerfilPersona.objects.get(codigo="dirigente_social"))
    fusionar(a, b, None, "misma persona")
    b.refresh_from_db()
    assert b.fecha_nacimiento == date(1990, 1, 1) and b.perfiles.filter(codigo="dirigente_social").exists()


def test_exportar_datos_personales_incluye_lo_nuevo(cg, persona):
    from io import BytesIO

    from openpyxl import load_workbook

    Consentimiento.objects.create(persona=persona, tipo="datos")
    r = cg.get(f"/personas/{persona.pk}/datos-personales")
    libro = load_workbook(BytesIO(r.content))
    assert "Consentimientos" in libro.sheetnames and "Comunidades" in libro.sheetnames
