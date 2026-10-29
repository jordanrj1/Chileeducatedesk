"""Vacíos detectados en la revisión de dudas: caducidad, lista de espera, constancia y derecho de acceso."""
from datetime import date, time, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.actividades import servicios
from apps.actividades.models import Participacion
from apps.actividades.reglas import revisar
from apps.nucleo.models import Auditoria

pytestmark = pytest.mark.django_db


def test_panel_avisa_lo_que_vence(cg, actividad, persona, universidad):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from apps.documentos.models import Documento
    from apps.nucleo.models import TipoDocumento
    from apps.personas.models import VerificacionPersona

    hoy = date.today()
    Documento.objects.create(nombre="Certificado de directorio", archivo=SimpleUploadedFile("c.pdf", b"%PDF"), tipo=TipoDocumento.objects.get(codigo="certificado"), organizacion=universidad, fecha_vencimiento=hoy + timedelta(days=5))
    Documento.objects.create(nombre="Lejano", archivo=SimpleUploadedFile("l.pdf", b"%PDF"), tipo=TipoDocumento.objects.get(codigo="otro"), actividad=actividad, fecha_vencimiento=hoy + timedelta(days=200))
    VerificacionPersona.objects.create(persona=persona, tipo="Certificado de antecedentes", fecha=hoy - timedelta(days=300), vigente_hasta=hoy - timedelta(days=1))
    universidad.vigencia_hasta = hoy + timedelta(days=20)
    universidad.save()
    html = cg.get("/dashboard").content.decode()
    assert "Datos por vencer" in html and "Certificado de directorio" in html and "Lejano" not in html
    assert "Certificado de antecedentes" in html and "Vigencia legal de Universidad Ficticia" in html


def test_lista_de_espera_no_ocupa_cupo(cg, actividad, lugar, gestor, crear_persona, rol, estado_part):
    actividad.responsable = crear_persona()
    actividad.save()
    servicios.programar(actividad, gestor, date.today() + timedelta(days=5), time(10), time(11), lugar, None)
    for _ in range(20):
        Participacion.objects.create(actividad=actividad, persona=crear_persona(), rol=rol("participante"), estado=estado_part("confirmado"))
    espera = Participacion.objects.create(actividad=actividad, persona=crear_persona(), rol=rol("participante"), estado=estado_part("lista_espera"))
    assert not espera.activa and espera.en_espera
    assert "capacidad" not in {p.codigo for p in revisar(actividad)}
    html = cg.get(f"/actividades/{actividad.pk}?tab=beneficiarios").content.decode()  # F26 P3a: los cupos van con los beneficiarios
    assert "Cupos ocupados: 20 de 20" in html and "en lista de espera: 1" in html and "sin cupos" in html
    # al confirmarla desde la espera, vuelve a contar y aparece el exceso
    espera.estado = estado_part("confirmado")
    espera.save()
    assert "capacidad" in {p.codigo for p in revisar(actividad)}


def test_constancia_solo_con_asistencia(cg, actividad, persona, rol, estado_part):
    p = Participacion.objects.create(actividad=actividad, persona=persona, rol=rol("voluntario"), estado=estado_part("confirmado"))
    assert cg.get(f"/actividades/{actividad.pk}/participantes/{p.pk}/constancia").status_code == 404
    p.asistencia = True
    p.save()
    html = cg.get(f"/actividades/{actividad.pk}/participantes/{p.pk}/constancia").content.decode()
    assert "Constancia de participación" in html and "Ana María Pérez Soto" in html and "voluntario" in html and "Imprimir" in html
    assert Auditoria.objects.filter(accion="constancia_emitida").exists()
    assert "/constancia" in cg.get(f"/actividades/{actividad.pk}?tab=asistencia").content.decode()


def test_constancia_de_otra_actividad_no(cg, actividad, crear_actividad, persona, rol, estado_part):
    otra = crear_actividad("Otra")
    p = Participacion.objects.create(actividad=otra, persona=persona, rol=rol("voluntario"), estado=estado_part("asistio"), asistencia=True)
    assert cg.get(f"/actividades/{actividad.pk}/participantes/{p.pk}/constancia").status_code == 404


def test_exportar_datos_personales(cg, cs, persona, actividad, rol, estado_part):
    from apps.nucleo.models import TipoContacto
    from apps.personas.models import Contacto, ContactoEmergencia

    Contacto.objects.create(persona=persona, tipo=TipoContacto.objects.get(codigo="correo"), valor="ana@ejemplo.invalid")
    ContactoEmergencia.objects.create(persona=persona, nombre="Madre Ficticia", telefono="+56911112222")
    Participacion.objects.create(actividad=actividad, persona=persona, rol=rol("voluntario"), estado=estado_part("confirmado"))
    r = cg.get(f"/personas/{persona.pk}/datos-personales")
    libro = load_workbook(BytesIO(r.content))
    assert r.status_code == 200 and {"Datos", "Contactos", "Participaciones", "Emergencia"} <= set(libro.sheetnames)
    assert any("ana@ejemplo.invalid" in str(c) for fila in libro["Contactos"].iter_rows(values_only=True) for c in fila)
    assert Auditoria.objects.filter(accion="datos_personales_exportados").exists()
    assert cs.get(f"/personas/{persona.pk}/datos-personales").status_code == 403


def test_exportar_datos_sin_permiso_de_sensibles_omite_emergencia(cg, gestor, persona):
    from django.contrib.auth.models import Permission

    gestor.groups.first().permissions.remove(Permission.objects.get(codename="ver_sensibles"))
    c = type(cg)()
    c.force_login(type(gestor).objects.get(pk=gestor.pk))
    libro = load_workbook(BytesIO(c.get(f"/personas/{persona.pk}/datos-personales").content))
    assert "Emergencia" not in libro.sheetnames and "Verificaciones" not in libro.sheetnames
